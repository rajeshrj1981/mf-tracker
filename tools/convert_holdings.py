#!/usr/bin/env python3
"""Convert Holdings.xlsx into seed JSON for the MF Tracker app.

Usage: python3 tools/convert_holdings.py Holdings.xlsx src/main/resources/seed/holdings.json
Also prints a data-quality report (also written next to the JSON as holdings-warnings.txt).
"""
import json, re, sys
from datetime import datetime
from openpyxl import load_workbook

CATEGORY = {
    "Small Cap Fund": "Small Cap", "Large Cap": "Large Cap", "Large&Mid Cap": "Large & Mid Cap",
    "Mid Cap": "Mid Cap", "Multi Cap": "Multi Cap", "Flexi Cap": "Flexi Cap",
    "Balance Advantage": "Balance Advantage", "Thematic": "Thematic",
    "Multi Asset Value": "Multi Asset & Value", "Debt Funds": "Debt",
}
HOUSES = [  # (prefix regex, fund house)
    (r"^bandhan", "Bandhan"), (r"^(bank of india|boa)\b", "Bank of India"), (r"^nippon", "Nippon India"),
    (r"^inve[sc]o", "Invesco"), (r"^sbi", "SBI"), (r"^motilal", "Motilal Oswal"), (r"^hsbc", "HSBC"),
    (r"^icici", "ICICI Prudential"), (r"^hdfc", "HDFC"), (r"^quant", "Quant"), (r"^dsp", "DSP"),
    (r"^franklin", "Franklin Templeton"), (r"^axis", "Axis"), (r"^lic", "LIC"), (r"^parakh|^parag", "PPFAS (check)"),
]
TYPOS = [(r"(?i)^inveso", "Invesco"), (r"(?i)enengy", "Energy")]

def clean_name(n):
    n = re.sub(r"\s+", " ", n).strip()
    for p, r in TYPOS: n = re.sub(p, r, n)
    return n

def house_of(n):
    for p, h in HOUSES:
        if re.search(p, n.lower()): return h
    return "Unknown"

def plan_of(name, brokers):
    l = name.lower()
    if "direct" in l: return "DIRECT"
    if "regular" in l: return "REGULAR"
    b = [x for x in brokers if x]
    if b and all(x == "Direct" for x in b): return "DIRECT"
    if b and all(x != "Direct" for x in b): return "REGULAR"   # e.g. bought via ICICI Direct platform
    return "UNKNOWN"

def num(v):
    return float(v) if isinstance(v, (int, float)) else None

def main(src, dst):
    wb = load_workbook(src, data_only=True)
    funds, warns = [], []
    for ws in wb.worksheets:
        cat = CATEGORY.get(ws.title.strip()) or CATEGORY.get(ws.title) or ws.title.strip()
        cur = None
        for row in ws.iter_rows(values_only=True):
            a, b, c, d, e, f = (list(row) + [None] * 6)[:6]
            if isinstance(a, datetime):                       # purchase row
                if cur is None: continue
                cur["purchases"].append({"date": a.date().isoformat(), "amount": num(b), "sheetNav": num(c),
                                         "units": num(d), "broker": (f or None)})
            elif isinstance(a, str) and a.strip() and not a.startswith(("Total", "Scheme Name", "Date of")) \
                    and isinstance(b, (int, float)):           # scheme header row
                cur = {"schemeName": clean_name(a), "rawName": a.strip(), "category": cat, "sheet": ws.title.strip(),
                       "sheetNav": num(b), "sheetNavDate": c.date().isoformat() if isinstance(c, datetime) else None,
                       "purchases": []}
                funds.append(cur)
    # derive plan / house, validate purchases
    sig = {}
    for fd in funds:
        fd["fundHouse"] = house_of(fd["schemeName"])
        fd["plan"] = plan_of(fd["schemeName"], [p["broker"] for p in fd["purchases"]])
        label = f'{fd["category"]} / {fd["schemeName"]} ({fd["plan"]})'
        seen = set()
        for p in fd["purchases"]:
            amt, units, snav = p["amount"], p["units"], p["sheetNav"]
            if not amt or not units:
                warns.append(f"{label}: {p['date']} missing amount/units - skipped"); p["skip"] = True; continue
            dnav = amt / units
            p["nav"] = round(dnav, 4)
            p["remark"] = None
            if snav is None:
                p["remark"] = "NAV missing in sheet; derived from amount/units"
            elif abs(snav - dnav) / dnav > 0.01:
                p["remark"] = f"Sheet NAV {snav} disagrees with amount/units ({dnav:.4f}); using amount/units"
            elif fd["sheetNav"] and abs(dnav / fd["sheetNav"] - 1) > 0.35:
                p["remark"] = f"Purchase NAV {dnav:.2f} is far from current NAV {fd['sheetNav']} - please verify"
            if fd["sheetNav"] and abs(dnav / fd["sheetNav"] - 1) > 0.35 and not p["remark"]:
                pass
            if p["remark"]: warns.append(f"{label}: {p['date']} Rs{amt:,.0f} - {p['remark']}")
            k = (p["date"], amt, units)
            if k in seen:
                p["remark"] = ((p["remark"] + "; ") if p["remark"] else "") + "Possible duplicate row"
                warns.append(f"{label}: {p['date']} Rs{amt:,.0f} - possible duplicate row (same date, amount, units)")
            seen.add(k)
        fd["purchases"] = [p for p in fd["purchases"] if not p.get("skip")]
        s = tuple((p["date"], p["amount"], p["units"]) for p in fd["purchases"])
        if s: sig.setdefault(s, []).append(label)
    for s, labels in sig.items():
        if len(labels) > 1 and len(s) > 1:
            warns.append("IDENTICAL purchase lists (likely copy-paste): " + "  <->  ".join(labels))
    for fd in funds:
        if not fd["sheetNavDate"]: warns.append(f'{fd["schemeName"]}: no NAV date in sheet (NAV treated as stale)')
        fd["purchaseCount"] = len(fd["purchases"])
    json.dump({"funds": funds}, open(dst, "w"), indent=1)
    rep = dst.rsplit(".", 1)[0] + "-warnings.txt"
    open(rep, "w").write("\n".join(warns) + "\n")
    inv = sum(p["amount"] for fd in funds for p in fd["purchases"])
    print(f"{len(funds)} fund blocks, {sum(len(f['purchases']) for f in funds)} purchases, invested Rs{inv:,.0f}")
    print("\n".join(warns))

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
